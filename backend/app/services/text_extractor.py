import io
import base64
from typing import Tuple
from PIL import Image
import pypdf
import docx

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts clean text content from PDF binary stream."""
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages_text.append(text)
        return "\n\n".join(pages_text).strip()
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts clean text content from DOCX binary stream."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs_text.append(" | ".join(row_text))
        return "\n\n".join(paragraphs_text).strip()
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extracts text from plain text files with automatic encoding detection."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            decoded = file_bytes.decode(encoding).strip()
            if decoded:
                return decoded
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace").strip()

def extract_text_from_image(file_bytes: bytes) -> Tuple[str, bool]:
    """
    Validates image and extracts text via OCR.
    Returns (extracted_text, has_ocr_text).
    If OCR engine is not locally present or returns empty, returns a base64 marker
    so downstream multimodal vision models can process it directly.
    """
    try:
        img = Image.open(io.BytesIO(file_bytes))
        img.verify()
        # Re-open after verify() closes the stream
        img = Image.open(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Corrupted or invalid image file: {str(e)}")

    # Attempt pytesseract if tesseract is installed
    try:
        import pytesseract
        ocr_text = pytesseract.image_to_string(img)
        if ocr_text and ocr_text.strip():
            return ocr_text.strip(), True
    except Exception:
        pass

    # Fallback to vision payload for multimodal LLM processing
    b64_data = base64.b64encode(file_bytes).decode("utf-8")
    return f"__VISION_IMAGE_B64__:{img.format or 'JPEG'}:{b64_data}", False

def extract_resume_text(filename: str, file_bytes: bytes) -> Tuple[str, str]:
    """
    Detects file type by extension and extracts readable text or vision payload.
    Supports: .pdf, .docx, .txt, .png, .jpg, .jpeg.
    Returns (extracted_text, file_type).
    """
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes), "pdf"
    elif lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes), "docx"
    elif lower.endswith(".txt"):
        return extract_text_from_txt(file_bytes), "txt"
    elif lower.endswith(".png"):
        text, _ = extract_text_from_image(file_bytes)
        return text, "png"
    elif lower.endswith(".jpg") or lower.endswith(".jpeg"):
        text, _ = extract_text_from_image(file_bytes)
        return text, "jpeg"
    else:
        raise ValueError(
            "Unsupported resume format. Supported formats: PDF, DOCX, TXT, PNG, JPEG."
        )
